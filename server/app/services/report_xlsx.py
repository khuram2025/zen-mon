"""Excel (.xlsx) output for section-engine reports.

``render_xlsx(meta, sections)`` turns the neutral section model (see
report_sections) into a workbook a reader can sort, filter and reuse:

- **Cover** — organisation name and logo (PNG/JPEG), title, category,
  description, period with its timezone label, generated time, scope,
  classification, footer text and disclaimer;
- **Summary** — every headline figure (section, KPI, value, detail), plus a row
  for each section that was withheld (permissions) or failed;
- one numbered sheet per section with data: each table as an Excel table
  (filterable, brand-coloured header), chart series as "chart data" tables,
  and the section's notes.

Cells holding numbers are written as numbers: percentages as fractions with a
percent format, figures with a unit keep the unit in the number format so they
display as before but sort and sum as numbers. Text that a spreadsheet would
read as a formula (leading ``= + - @``, tab or CR) is prefixed with an
apostrophe. Output depends only on its inputs (no clock reads) and needs no
database access.
"""

from __future__ import annotations

import io
import logging
import math
import re
import sys
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.services import report_colors as rc
from app.services.branding import Branding, sniff_logo

logger = logging.getLogger("zenplus.reports")

FONT = "Calibri"
MAX_CELL_CHARS = 32767
MAX_CHART_POINTS = 5000
NEUTRAL_VALUES = {"", "-", "—", "–", "n/a", "na", "none", "unknown", "null"}
NUMERIC_STYLES = {"num", "pct", "pct-bar"}
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\\x00-\x1f]")
_HEADER_TEXT_HINTS = {"status", "state", "health", "severity", "ip", "ip address",
                      "address", "mac", "port", "version", "id", "serial"}
_NUM_RE = re.compile(
    r"^(?P<sign>[-+−]?)(?P<int>\d{1,3}(?:,\d{3})+|\d+)(?P<frac>\.\d+)?"
    r"(?:\s*(?P<unit>%|[A-Za-zµ]{1,8}(?:/[A-Za-z]{1,4})?))?$"
)

MUTED = "6B7280"
BODY = "1F2937"

EXCEL_DATE_FORMATS = {
    "YYYY-MM-DD": "yyyy-mm-dd",
    "DD/MM/YYYY": "dd/mm/yyyy",
    "MM/DD/YYYY": "mm/dd/yyyy",
    "DD.MM.YYYY": "dd.mm.yyyy",
    "DD-MMM-YYYY": "dd-mmm-yyyy",
    "MMM D, YYYY": "mmm d, yyyy",
}


# ─── Cell values ────────────────────────────────────────────────────────────

def safe_text(value: Any) -> str:
    """Spreadsheet-safe text: XML-illegal control characters removed, capped
    at Excel's cell limit, and formula triggers neutralised with a leading
    apostrophe (CSV/formula injection)."""
    s = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(s) > MAX_CELL_CHARS - 1:
        s = s[:MAX_CELL_CHARS - 2] + "…"
    if s.startswith(FORMULA_TRIGGERS):
        s = "'" + s
    return s


@dataclass(frozen=True)
class ParsedNumber:
    value: float | int
    kind: str          # "num" | "pct"
    unit: str          # "" for a bare number
    decimals: int
    grouped: bool      # had thousands separators


def parse_number(text: str) -> Optional[ParsedNumber]:
    """'99.95%', '1,234', '12.5 ms', '-3' → ParsedNumber; anything else
    (ranges, '<0.01%', durations like '5h 2m', leading zeros) → None."""
    m = _NUM_RE.match(text.strip())
    if not m:
        return None
    int_part = m.group("int")
    if len(int_part) > 1 and int_part.startswith("0"):
        return None  # identifiers such as '007' keep their leading zero
    frac = m.group("frac") or ""
    neg = m.group("sign") in ("-", "−")
    digits = int_part.replace(",", "")
    try:
        value: float | int = float(digits + frac) if frac else int(digits)
    except ValueError:
        return None
    if neg:
        value = -value
    unit = m.group("unit") or ""
    decimals = min(len(frac) - 1, 6) if frac else 0
    if unit == "%":
        return ParsedNumber(value / 100, "pct", "%", decimals, "," in int_part)
    return ParsedNumber(value, "num", unit, decimals, "," in int_part)


def number_format(p: ParsedNumber) -> str:
    if p.kind == "pct":
        return "0." + "0" * max(2, p.decimals) + "%"
    base = "#,##0" if p.grouped else "0"
    if p.decimals:
        base += "." + "0" * p.decimals
    if p.unit:
        base += f'" {p.unit}"'
    return base


def excel_datetime_format(branding: Optional[Branding], seconds: bool = False) -> str:
    """Excel number format matching the organisation's date/time format."""
    b = branding or Branding()
    date = EXCEL_DATE_FORMATS.get(b.date_format, "yyyy-mm-dd")
    if b.time_format == "12h":
        time = "h:mm:ss AM/PM" if seconds else "h:mm AM/PM"
    else:
        time = "hh:mm:ss" if seconds else "hh:mm"
    return f"{date} {time}"


def _parse_ts(value: Any) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:  # epoch seconds or milliseconds
            v = float(value)
            return datetime.fromtimestamp(v / 1000 if v > 1e11 else v, tz=timezone.utc)
        except (OverflowError, OSError, ValueError):
            return None
    if isinstance(value, str) and value:
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
    return None


# ─── Rendering context ──────────────────────────────────────────────────────

@dataclass
class _Ctx:
    meta: dict
    branding: Branding
    brand: str
    header_fill: PatternFill
    header_font: Font
    title_color: str
    datetime_format: str
    tz_label: str
    company: str


def _context(meta: dict) -> _Ctx:
    b = meta.get("branding")
    branding = b if isinstance(b, Branding) else Branding()
    brand = rc.normalize(meta.get("brand_color") or branding.brand_color)
    return _Ctx(
        meta=meta,
        branding=branding,
        brand=brand,
        header_fill=PatternFill("solid", start_color=brand[1:], end_color=brand[1:]),
        header_font=Font(name=FONT, bold=True, color=rc.text_on(brand)[1:], size=10.5),
        title_color=rc.ink(brand, 3.0)[1:],
        datetime_format=excel_datetime_format(branding),
        tz_label=str(meta.get("tz_label") or branding.tz_label()),
        company=str(meta.get("company_name") or branding.display_name),
    )


class _Widths:
    """Column widths from the content actually written (header and cells)."""

    def __init__(self, fixed: Optional[dict[int, float]] = None) -> None:
        self.chars: dict[int, int] = {}
        self.fixed: dict[int, float] = dict(fixed or {})

    def see(self, col: int, text: str, extra: int = 0) -> None:
        longest = max((len(line) for line in text.split("\n")), default=0) + extra
        if longest > self.chars.get(col, 0):
            self.chars[col] = longest

    def apply(self, ws, minimum: int = 8, maximum: int = 60) -> None:
        for col, n in self.chars.items():
            ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(n + 2, maximum))

    def width(self, col: int, default: float = 8.43) -> float:
        if col in self.fixed:
            return self.fixed[col]
        n = self.chars.get(col)
        return default if n is None else max(8, min(n + 2, 60))


RIGHT = Alignment(horizontal="right")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
HEADER_ALIGN = Alignment(vertical="center")
BODY_FONT = Font(name=FONT, size=10.5, color=BODY)
MUTED_ITALIC_FONT = Font(name=FONT, size=10, italic=True, color=MUTED)
TABLE_TITLE_FONT = Font(name=FONT, bold=True, size=11.5, color=BODY)


def _put(ws, row: int, col: int, value: Any, *, font: Optional[Font] = None,
         fmt: Optional[str] = None, fill: Optional[PatternFill] = None,
         align: Optional[Alignment] = None):
    """Write a cell. Assigning a style makes openpyxl hash the style objects,
    which dominates large workbooks, so the resulting style array is cached
    per workbook (keyed by the identity of shared style objects) and copied
    onto later unstyled cells with the same combination."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = safe_text(value)
    else:
        cell.value = value
    if font is None and fmt is None and fill is None and align is None:
        return cell
    wb = ws.parent
    cache = getattr(wb, "_report_style_cache", None)
    if cache is None:
        cache = {}
        wb._report_style_cache = cache
    key = (id(font), fmt, id(fill), id(align))
    hit = cache.get(key)
    fresh = not cell.has_style
    if hit is not None and fresh:
        cell._style = copy(hit[0])
        return cell
    if font is not None:
        cell.font = font
    if fmt is not None:
        cell.number_format = fmt
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if hit is None and fresh:
        # Keep the style objects alive so their ids cannot be reused.
        cache[key] = (copy(cell._style), (font, fill, align))
    return cell


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:,.2f}"
    return str(value)


def _wrapped_row(ws, row: int, text: str, first_col: int, last_col: int, widths: _Widths,
                 font: Font) -> None:
    """A merged, wrapped text row; Excel does not auto-fit merged rows, so
    the height is estimated from the text length and merged width."""
    _put(ws, row, first_col, text, font=font,
         align=WRAP_TOP)
    if last_col > first_col:
        ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    span = sum(widths.width(c) for c in range(first_col, last_col + 1))
    per_line = max(int(span * 1.15), 20)
    lines = sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))
    if lines > 1:
        ws.row_dimensions[row].height = min(15.0 * lines + 2, 409)


def _hf(text: str, limit: int = 70) -> str:
    """Header/footer text: '&' is a control character there."""
    t = ILLEGAL_CHARACTERS_RE.sub("", str(text or "")).replace("\n", " ")
    if len(t) > limit:
        t = t[:limit - 1] + "…"
    return t.replace("&", "&&")


def _print_setup(ws, ctx: _Ctx, *, landscape: bool) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER if ctx.meta.get("paper_size") == "Letter" \
        else ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.oddHeader.left.text = _hf(ctx.company)
    classification = str(ctx.meta.get("confidentiality") or "")
    if classification:
        ws.oddHeader.right.text = _hf(classification.upper(), 30)
    ws.oddFooter.left.text = _hf(ctx.meta.get("footer_text") or ctx.company)
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = _hf(ctx.meta.get("title") or "Report", 50)


# ─── Sheet names ────────────────────────────────────────────────────────────

def sheet_name(number: int, title: str, used: set[str], digits: int = 2) -> str:
    """'03 Interface Utilization' — sanitised, ≤31 chars, unique
    (case-insensitively, as Excel compares them)."""
    clean = _SHEET_FORBIDDEN.sub(" ", str(title or "")).strip().strip("'")
    clean = re.sub(r"\s+", " ", clean) or "Section"
    prefix = f"{number:0{digits}d} "
    name = (prefix + clean)[:31].rstrip().rstrip("'")
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = name[:31 - len(suffix)].rstrip() + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


LINK_FONT = Font(name=FONT, color="1D4ED8", underline="single", size=10.5)


def _internal_link(cell, sheet: str) -> None:
    """Hyperlink to cell A1 of another sheet of this workbook."""
    quoted = sheet.replace("'", "''")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{quoted}'!A1", display=str(cell.value))
    cell.font = LINK_FONT


# ─── Tables ─────────────────────────────────────────────────────────────────

def _column_plans(headers: list[str], rows: list[list], styles: list) -> list[Optional[set[str]]]:
    """Per column: the units whose values convert to numbers, or None to keep
    the column as text. Styled num/pct/pct-bar columns convert where a value
    parses; unstyled columns convert only when every value parses. Mixed
    units convert only their bare numbers (512 Mbps and 1.2 Gbps must not be
    compared as 512 and 1.2)."""
    ncols = len(headers)
    plans: list[Optional[set[str]]] = []
    for ci in range(ncols):
        style = styles[ci] if ci < len(styles) else None
        if style and style not in NUMERIC_STYLES:
            plans.append(None)
            continue
        if not style and str(headers[ci]).strip().lower() in _HEADER_TEXT_HINTS:
            plans.append(None)
            continue
        units: set[str] = set()
        parsed = 0
        all_parse = True
        for r in rows:
            v = r[ci] if ci < len(r) else None
            if v is None or isinstance(v, bool):
                continue
            if isinstance(v, (int, float)):
                parsed += 1
                continue
            s = str(v).strip()
            if s.lower() in NEUTRAL_VALUES:
                continue
            p = parse_number(s)
            if p is None:
                all_parse = False
                continue
            parsed += 1
            if p.kind == "num":
                units.add(p.unit)
        if not style and (not all_parse or not parsed):
            plans.append(None)
            continue
        allowed = {"", "%"} | (units if len(units - {""}) <= 1 else set())
        plans.append(allowed)
    return plans


def _unique_headers(headers: list[str], ncols: int) -> list[str]:
    """Excel table column names: non-empty, single-line, unique ignoring case."""
    out: list[str] = []
    seen: set[str] = set()
    for i in range(ncols):
        raw = headers[i] if i < len(headers) else ""
        name = re.sub(r"\s+", " ", ILLEGAL_CHARACTERS_RE.sub("", str(raw or ""))).strip()[:200]
        name = name or f"Column {i + 1}"
        cand, n = name, 2
        while safe_text(cand).lower() in seen:
            cand = f"{name} {n}"
            n += 1
        seen.add(safe_text(cand).lower())
        out.append(cand)
    return out


def _write_value(ws, row: int, col: int, value: Any, plan: Optional[set[str]],
                 ctx: _Ctx, widths: _Widths, font: Font) -> None:
    if value is None:
        return
    if isinstance(value, bool):
        value = "Yes" if value else "No"
    if isinstance(value, datetime):
        local = ctx.branding.localize(value).replace(tzinfo=None)
        _put(ws, row, col, local, font=font, fmt=ctx.datetime_format)
        widths.see(col, "0000-00-00 00:00 PM")
        return
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            _put(ws, row, col, "—", font=font)
            return
        _put(ws, row, col, value, font=font)
        widths.see(col, _display(value))
        return
    text = str(value)
    if plan is not None:
        p = parse_number(text) if text.strip().lower() not in NEUTRAL_VALUES else None
        if p is not None and p.unit in plan:
            _put(ws, row, col, p.value, font=font, fmt=number_format(p),
                 align=RIGHT)
            widths.see(col, text.strip())
            return
        _put(ws, row, col, text, font=font, align=RIGHT)
        widths.see(col, text)
        return
    _put(ws, row, col, text, font=font)
    widths.see(col, text)


class _SheetWriter:
    def __init__(self, ws, ctx: _Ctx, sheet_no: int) -> None:
        self.ws = ws
        self.ctx = ctx
        self.sheet_no = sheet_no
        self.widths = _Widths()
        self.row = 1
        self.table_count = 0
        self.first_header_row: Optional[int] = None
        self.body_font = BODY_FONT
        self.muted_font = MUTED_ITALIC_FONT

    def table(self, title: str, headers: list, rows: list[list], styles: Optional[list] = None,
              plans: Optional[list] = None) -> None:
        ws, ctx = self.ws, self.ctx
        headers = [("" if h is None else str(h)) for h in (headers or [])]
        rows = [list(r) if isinstance(r, (list, tuple)) else [r] for r in (rows or [])]
        ncols = max([len(headers)] + [len(r) for r in rows] + [1])
        names = _unique_headers(headers, ncols)
        if plans is None:
            plans = _column_plans(names, rows, list(styles or []))

        _put(ws, self.row, 1, title, font=TABLE_TITLE_FONT)
        self.row += 1
        header_row = self.row
        for ci, name in enumerate(names, start=1):
            _put(ws, header_row, ci, name, font=ctx.header_font, fill=ctx.header_fill,
                 align=HEADER_ALIGN)
            self.widths.see(ci, name, extra=2)
        if self.first_header_row is None:
            self.first_header_row = header_row
        self.row += 1

        if not rows:
            _put(ws, self.row, 1, "No data for this window", font=self.muted_font)
            self.row += 2
            return

        first_data = self.row
        for r in rows:
            for ci in range(ncols):
                v = r[ci] if ci < len(r) else None
                _write_value(ws, self.row, ci + 1, v, plans[ci] if ci < len(plans) else None,
                             ctx, self.widths, self.body_font)
            self.row += 1
        last_data = self.row - 1

        self.table_count += 1
        ref = f"A{header_row}:{get_column_letter(ncols)}{last_data}"
        tab = Table(displayName=f"Table{self.sheet_no:02d}_{self.table_count}", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True,
                                            showColumnStripes=False, showFirstColumn=False,
                                            showLastColumn=False)
        ws.add_table(tab)

        styles = list(styles or [])
        for ci in range(ncols):
            if ci < len(styles) and styles[ci] == "pct-bar" and plans[ci] is not None:
                col = get_column_letter(ci + 1)
                ws.conditional_formatting.add(
                    f"{col}{first_data}:{col}{last_data}",
                    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                color=self.ctx.brand[1:], showValue=True))
        self.row += 1  # blank line between tables

    def finish(self) -> None:
        self.widths.apply(self.ws)
        if self.first_header_row is not None:
            self.ws.freeze_panes = f"A{self.first_header_row + 1}"
            self.ws.print_title_rows = f"{self.first_header_row}:{self.first_header_row}"


def _chart_tables(chart: dict, ctx: _Ctx, include_categorical: bool) -> Optional[tuple]:
    """(title, headers, rows) for a chart's native series, or None."""
    spec = chart.get("series")
    if not isinstance(spec, dict):
        return None
    unit = str(spec.get("unit") or "").strip()
    suffix = f" ({unit})" if unit else ""
    title = f"Chart data — {chart.get('title') or 'Chart'}"
    lines = spec.get("series")
    if isinstance(lines, list) and lines:
        series = [(str(s.get("name") or f"Series {i + 1}"), s.get("points") or [])
                  for i, s in enumerate(lines) if isinstance(s, dict)]
    else:
        series = [(str(chart.get("title") or "Value"), spec.get("points") or [])]
    series = [(n, [p for p in pts if isinstance(p, dict)][:MAX_CHART_POINTS]) for n, pts in series]
    if not any(pts for _, pts in series):
        return None

    if all("t" in p for _, pts in series for p in pts):
        keys: list = []
        table: dict = {}
        for si, (_, pts) in enumerate(series):
            for p in pts:
                k = p.get("t")
                if k not in table:
                    keys.append(k)
                    table[k] = [None] * len(series)
                table[k][si] = p.get("v")
        parsed = {k: _parse_ts(k) for k in keys}
        if all(parsed[k] is not None for k in keys):
            keys.sort(key=lambda k: ctx.branding.localize(parsed[k]))
        rows = [[parsed[k] if parsed[k] is not None else str(k)] + table[k] for k in keys]
        headers = [f"Time ({ctx.tz_label})"] + [f"{n}{suffix}" for n, _ in series]
        return title, headers, rows

    if not include_categorical:
        return None
    rows = []
    for _, pts in series[:1]:
        for p in pts:
            rows.append([str(p.get("label") if p.get("label") is not None else p.get("t", "")),
                         p.get("value", p.get("v"))])
    return title, ["Label", f"Value{suffix}"], rows


def _numeric_plan_for_chart(headers: list) -> list:
    return [None] + [{"", "%"}] * (len(headers) - 1)


# ─── Sheets ─────────────────────────────────────────────────────────────────

def _logo_image(meta: dict):
    """openpyxl image for a PNG/JPEG light-background logo, scaled to fit."""
    data = meta.get("logo_bytes")
    if not data:
        return None
    mime = sniff_logo(bytes(data))
    if mime not in ("image/png", "image/jpeg"):
        return None
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(io.BytesIO(bytes(data)))
        w, h = float(img.width or 0), float(img.height or 0)
        if w <= 0 or h <= 0:
            return None
        scale = min(1.0, 56.0 / h, 240.0 / w)
        img.width, img.height = max(1, round(w * scale)), max(1, round(h * scale))
        return img
    except Exception:
        return None


def _cover(ws, ctx: _Ctx, sections: list[dict]) -> None:
    meta = ctx.meta
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    widths = _Widths(fixed={1: 22, 2: 90})
    row = 1

    classification = str(meta.get("confidentiality") or "").strip()
    if classification:
        bg = rc.classification_bg(classification)
        fill = PatternFill("solid", start_color=bg[1:], end_color=bg[1:])
        for col in (1, 2):
            ws.cell(row=row, column=col).fill = fill
        _put(ws, row, 1, classification.upper(),
             font=Font(name=FONT, bold=True, size=11, color=rc.text_on(bg)[1:]),
             align=Alignment(horizontal="center", vertical="center"))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 22
        row += 2

    img = _logo_image(meta)
    if img is not None:
        ws.add_image(img, f"A{row}")
        ws.row_dimensions[row].height = img.height * 0.75 + 6
        row += 1

    _put(ws, row, 1, ctx.company, font=Font(name=FONT, bold=True, size=16, color=ctx.title_color))
    row += 1
    brand_fill = PatternFill("solid", start_color=ctx.brand[1:], end_color=ctx.brand[1:])
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = brand_fill
    ws.row_dimensions[row].height = 4
    row += 2

    _put(ws, row, 1, str(meta.get("title") or "Report"),
         font=Font(name=FONT, bold=True, size=20, color=BODY))
    ws.row_dimensions[row].height = 28
    row += 1
    if meta.get("description"):
        _wrapped_row(ws, row, str(meta["description"]), 1, 2, widths,
                     Font(name=FONT, size=11, color=MUTED))
        row += 1
    row += 1

    label_font = Font(name=FONT, bold=True, size=10.5, color=MUTED)
    value_font = Font(name=FONT, size=11, color=BODY)
    tz = str(meta.get("timezone") or ctx.branding.timezone)
    facts = [
        ("Category", meta.get("category")),
        ("Period", meta.get("period_label")),
        ("Timezone", f"{tz} ({ctx.tz_label})" if tz and tz not in ("UTC", ctx.tz_label) else ctx.tz_label),
        ("Generated", meta.get("generated_label")),
        ("Scope", meta.get("scope_label")),
        ("Classification", classification),
        ("Prepared by", ctx.company),
        ("Contact", meta.get("contact_line")),
        ("Address", meta.get("company_address")),
        ("Sections", str(len(sections)) if sections else ""),
    ]
    top = Alignment(vertical="top")
    for label, value in facts:
        if not value:
            continue
        _put(ws, row, 1, label, font=label_font, align=top)
        _put(ws, row, 2, str(value), font=value_font, align=WRAP_TOP)
        lines = sum(max(1, math.ceil(len(part) / 100)) for part in str(value).split("\n"))
        if lines > 1:
            ws.row_dimensions[row].height = min(15.0 * lines + 2, 409)
        row += 1

    link = ws.cell(row=row + 1, column=1, value="Summary of figures →")
    _internal_link(link, "Summary")
    row += 3

    if meta.get("footer_text"):
        _wrapped_row(ws, row, str(meta["footer_text"]), 1, 2, widths,
                     Font(name=FONT, size=10, color=MUTED))
        row += 1
    if meta.get("disclaimer"):
        _wrapped_row(ws, row, str(meta["disclaimer"]), 1, 2, widths,
                     Font(name=FONT, size=9, italic=True, color=MUTED))
        row += 1
    _print_setup(ws, ctx, landscape=False)


def _restriction_text(section: dict) -> str:
    requires = section.get("requires")
    if not requires:
        rs = sys.modules.get("app.services.report_sections")
        entry = getattr(rs, "SECTION_REGISTRY", {}).get(section.get("section_id") or section.get("id")) \
            if rs is not None else None
        requires = (entry or {}).get("requires")
    if requires:
        need = ", ".join(str(r) for r in requires) if isinstance(requires, (list, tuple, set, frozenset)) \
            else str(requires)
        return f"Not included — requires the {need} permission"
    note = next(iter(section.get("notes") or []), "")
    m = re.search(r"requires the (.+?) permission", str(note))
    if m:
        return f"Not included — requires the {m.group(1)} permission"
    return "Not included"


def _summary(ws, ctx: _Ctx, sections: list[dict], sheet_for: dict[int, str]) -> None:
    meta = ctx.meta
    _put(ws, 1, 1, "Summary", font=Font(name=FONT, bold=True, size=16, color=ctx.title_color))
    _put(ws, 2, 1, " · ".join(str(p) for p in (meta.get("title"), meta.get("period_label")) if p),
         font=Font(name=FONT, size=10, color=MUTED))
    header_row = 4
    headers = ["Section", "KPI", "Value", "Detail"]
    for ci, h in enumerate(headers, start=1):
        _put(ws, header_row, ci, h, font=ctx.header_font, fill=ctx.header_fill)
    widths = _Widths()
    for ci, h in enumerate(headers, start=1):
        widths.see(ci, h, extra=2)
    muted = Font(name=FONT, size=10.5, italic=True, color=MUTED)
    warn = Font(name=FONT, size=10.5, italic=True, color=rc.ACCENT_INK["warning"][1:])
    sub_font = Font(name=FONT, size=10, color=MUTED)
    value_fonts: dict[str, Font] = {}

    def put(r: int, col: int, text: str, font: Font) -> None:
        if text:
            _put(ws, r, col, text, font=font)
            widths.see(col, text)

    row = header_row + 1
    for idx, s in enumerate(sections):
        title = str(s.get("title") or s.get("id") or "Section")
        sheet = sheet_for.get(idx)
        notes = [str(n) for n in (s.get("notes") or []) if n]
        if s.get("restricted"):
            text = _restriction_text(s)
            detail = notes[0] if notes and "requires the" not in notes[0] else ""
            entries = [(text, None, detail, muted)]
        elif s.get("error"):
            entries = [("Not generated — this section failed", None, notes[0] if notes else "", warn)]
        else:
            kpis = [k for k in (s.get("kpis") or []) if isinstance(k, dict)]
            entries = [(str(k.get("label") or ""), k, str(k.get("subtitle") or ""), BODY_FONT)
                       for k in kpis]
            if not entries:
                detail = f"Details on sheet “{sheet}”" if sheet else (notes[0] if notes else "")
                entries = [("—", None, detail, muted)]
        for label, kpi, detail, label_font in entries:
            put(row, 1, title, BODY_FONT)
            if sheet:
                _internal_link(ws.cell(row=row, column=1), sheet)
            put(row, 2, label, label_font)
            if kpi is not None:
                accent = str(kpi.get("accent") or "primary")
                font = value_fonts.get(accent)
                if font is None:
                    font = value_fonts[accent] = Font(name=FONT, size=10.5, bold=True,
                                                      color=rc.accent_ink(accent, ctx.brand)[1:])
                value = kpi.get("value")
                _write_value(ws, row, 3, value, {"", "%"} | _single_unit(value), ctx, widths, font)
            put(row, 4, detail, sub_font if kpi is not None else muted)
            row += 1
    last = max(row - 1, header_row)
    ws.auto_filter.ref = f"A{header_row}:D{last}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    widths.apply(ws, maximum=70)
    _print_setup(ws, ctx, landscape=True)


def _single_unit(value: Any) -> set[str]:
    p = parse_number(str(value)) if isinstance(value, str) else None
    return {p.unit} if p is not None else set()


def _section_sheet(ws, ctx: _Ctx, s: dict, sheet_no: int, tables: list, charts: list) -> None:
    w = _SheetWriter(ws, ctx, sheet_no)
    # Width hints first so wrapped description/notes rows estimate height
    # against the final column widths.
    for t in tables:
        for ci, h in enumerate(t.get("headers") or [], start=1):
            w.widths.see(ci, str(h), extra=2)
        for r in (t.get("rows") or [])[:500]:
            if isinstance(r, (list, tuple)):
                for ci, v in enumerate(r, start=1):
                    if isinstance(v, str):
                        w.widths.see(ci, v)
    span = max([len(t.get("headers") or []) for t in tables] +
               [len(c[1]) for c in charts] + [6])

    _put(ws, 1, 1, str(s.get("title") or "Section"),
         font=Font(name=FONT, bold=True, size=14, color=ctx.title_color))
    ws.row_dimensions[1].height = 20
    meta_line = " · ".join(str(p) for p in (s.get("category"), ctx.meta.get("period_label")) if p)
    w.row = 2
    if meta_line:
        _put(ws, 2, 1, meta_line, font=Font(name=FONT, size=9.5, color=MUTED))
        w.row = 3
    if s.get("description"):
        _wrapped_row(ws, w.row, str(s["description"]), 1, span, w.widths,
                     Font(name=FONT, size=10, color=MUTED))
        w.row += 1
    w.row += 1

    many = len(tables) + len(charts) > 1
    for i, t in enumerate(tables, start=1):
        title = str(t.get("title") or (s.get("title") if not many else f"Table {i}") or "Data")
        w.table(title, list(t.get("headers") or []), list(t.get("rows") or []),
                styles=list(t.get("styles") or []))
    for title, headers, rows in charts:
        w.table(title, headers, rows, plans=_numeric_plan_for_chart(headers))

    notes = [str(n) for n in (s.get("notes") or []) if n]
    if notes:
        _put(ws, w.row, 1, "Notes", font=Font(name=FONT, bold=True, size=11, color=BODY))
        w.row += 1
        for n in notes:
            color = rc.ACCENT_INK["warning"][1:] if "could not be generated" in n else MUTED
            _wrapped_row(ws, w.row, n, 1, span, w.widths,
                         Font(name=FONT, size=10, italic=True, color=color))
            w.row += 1
    w.finish()
    _print_setup(ws, ctx, landscape=True)


# ─── Entry point ────────────────────────────────────────────────────────────

def render_xlsx(meta: dict, sections: list[dict]) -> bytes:
    """Workbook bytes for a rendered report (Cover, Summary, one sheet per
    section with data)."""
    meta = meta or {}
    sections = [s for s in (sections or []) if isinstance(s, dict)]
    ctx = _context(meta)

    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    summary = wb.create_sheet("Summary")

    plans: list[tuple[int, dict, list, list]] = []
    for idx, s in enumerate(sections):
        if s.get("restricted") or s.get("error"):
            continue
        tables = [t for t in (s.get("tables") or []) if isinstance(t, dict)]
        charts = []
        for c in s.get("charts") or []:
            if not isinstance(c, dict):
                continue
            try:
                got = _chart_tables(c, ctx, include_categorical=not tables)
            except Exception:  # a malformed series must not cost the whole workbook
                logger.warning("chart data of section %s skipped", s.get("id"), exc_info=True)
                got = None
            if got:
                charts.append(got)
        if tables or charts:
            plans.append((idx, s, tables, charts))

    digits = 3 if len(sections) > 99 else 2
    used = {"cover", "summary"}
    sheet_for: dict[int, str] = {}
    for idx, s, _, _ in plans:
        sheet_for[idx] = sheet_name(idx + 1, str(s.get("title") or ""), used, digits)

    _cover(cover, ctx, sections)
    _summary(summary, ctx, sections, sheet_for)
    for idx, s, tables, charts in plans:
        ws = wb.create_sheet(sheet_for[idx])
        _section_sheet(ws, ctx, s, idx + 1, tables, charts)

    cover.sheet_properties.tabColor = ctx.brand[1:]
    summary.sheet_properties.tabColor = ctx.brand[1:]

    props = wb.properties
    props.title = safe_text(str(meta.get("title") or "Report"))[:255]
    props.creator = safe_text(ctx.company)[:255]
    props.lastModifiedBy = props.creator
    props.subject = safe_text(str(meta.get("category") or ""))[:255] or None
    props.description = safe_text(str(meta.get("description") or ""))[:1000] or None
    if meta.get("confidentiality"):
        props.keywords = safe_text(str(meta["confidentiality"]))
    generated = meta.get("generated_at")
    if isinstance(generated, datetime):
        stamp = generated.astimezone(timezone.utc).replace(tzinfo=None) if generated.tzinfo else generated
        props.created = stamp

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
