"""Spreadsheet and email outputs of reports.

- report_xlsx.render_xlsx: Cover/Summary/section sheets, typed numbers,
  formula-injection protection, logo, Excel tables, branding.
- report_sections email summary: email-safe HTML (no data: URIs, no SVG,
  bgcolor fallbacks, CID logo), size budget, classification banner, text
  alternative and the CID logo attachment.
- report_html (legacy email/page): organisation branding instead of the
  hard-coded product badge.
- export_service (legacy Excel/CSV): branding/timezone, injection
  protection, and the scope leak regression (an empty resolved device list
  must mean "no devices", never "every device").
"""

from __future__ import annotations

import asyncio
import base64
import csv
import io
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core import scoping  # noqa: E402
from app.services import export_service, report_html  # noqa: E402
from app.services import report_sections as rs  # noqa: E402
from app.services.branding import Branding, branding_from_values  # noqa: E402
from app.services.report_xlsx import parse_number, render_xlsx, safe_text, sheet_name  # noqa: E402

FRM = datetime(2026, 10, 1, tzinfo=timezone.utc)
TO = datetime(2026, 10, 8, tzinfo=timezone.utc)
SVG = b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10"/></svg>'


def png_bytes(w: int = 300, h: int = 100) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


def logo(data: bytes) -> dict:
    return {"data": base64.b64encode(data).decode(), "content_type": "image/png"}


def make_branding(**over) -> Branding:
    values = {"company_name": "Acme Networks", "timezone": "Asia/Riyadh", "brand_color": "#0B3D91",
              "report_footer": "Acme NOC", "report_disclaimer": "Internal use only.",
              "company_website": "https://acme.example", "confidentiality": "Confidential"}
    light = over.pop("logo", None)
    dark = over.pop("logo_dark", None)
    values.update(over)
    return branding_from_values(values, logo=light, logo_dark=dark)


def make_meta(branding: Branding | None = None, **over) -> dict:
    b = branding or make_branding()
    meta = asyncio.run(rs.build_report_meta(None, "Weekly Operations", FRM, TO,
                                            description="Weekly health of the estate.",
                                            category="Availability",
                                            scope_label="All monitored infrastructure",
                                            branding=b))
    meta.update(over)
    return meta


def sections() -> list[dict]:
    return [
        {"id": "device_availability", "title": "Device Availability", "category": "Availability",
         "description": "Per-device availability.",
         "kpis": [{"label": "Availability", "value": "99.95%", "accent": "success", "subtitle": "target 99.9%"},
                  {"label": "Outages", "value": "3", "accent": "danger"}],
         "charts": [{"title": "Trend", "series": {"kind": "area", "unit": "%", "points": [
             {"t": "2026-10-01T00:00:00+00:00", "v": 99.9}, {"t": "2026-10-02T00:00:00+00:00", "v": 100.0}]}}],
         "tables": [{"title": "Devices",
                     "headers": ["Device", "IP", "Status", "Uptime", "Avg RTT", "Samples"],
                     "styles": ["text", "mono", "status", "pct-bar", "num", "num"],
                     "rows": [["=HYPERLINK(\"http://evil\")", "10.0.0.1", "up", "99.95%", "12.5 ms", "1,234"],
                              ["@SUM(1+1)", "10.0.0.2", "down", "—", "—", "512"],
                              ["+cmd", "10.0.0.3", "up", "100%", "3.0 ms", "007"],
                              ["-2+3", "10.0.0.4", "up", "98.5%", "4.2 ms", "10"],
                              ["\tTabbed", "10.0.0.5", "up", "97%", "1.0 ms", "1"]]}],
         "notes": ["Two devices breached the availability target."]},
        {"id": "apm_services", "title": "APM Services", "restricted": True,
         "notes": ["Not included: viewing this section requires the apm.view permission."]},
        {"id": "top_outages", "title": "Top Outages", "error": True,
         "notes": ["This section could not be generated (ValueError). Other sections are unaffected."]},
        {"id": "traffic_kpis", "title": "Traffic Overview",
         "kpis": [{"label": "Peak", "value": "1.2 Gbps", "accent": "warning"}]},
    ]


def load(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data))


def cells(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


# ─── render_xlsx ────────────────────────────────────────────────────────────

def test_xlsx_has_cover_summary_and_section_sheets():
    wb = load(render_xlsx(make_meta(), sections()))
    assert wb.sheetnames == ["Cover", "Summary", "01 Device Availability"]
    assert wb.properties.title == "Weekly Operations"
    assert wb.properties.creator == "Acme Networks"

    cover = [v for row in cells(wb["Cover"]) for v in row if v is not None]
    assert "Acme Networks" in cover and "Weekly Operations" in cover
    assert "CONFIDENTIAL" in cover
    assert any(isinstance(v, str) and "GMT+3" in v and "2026-10-01 03:00" in v for v in cover)  # period + tz
    assert "Asia/Riyadh (GMT+3)" in cover
    assert "Availability" in cover and "Weekly health of the estate." in cover
    assert "All monitored infrastructure" in cover
    assert "Acme NOC" in cover and "Internal use only." in cover


def test_xlsx_summary_lists_kpis_restricted_and_failed_sections():
    ws = load(render_xlsx(make_meta(), sections()))["Summary"]
    rows = cells(ws)
    assert rows[3] == ["Section", "KPI", "Value", "Detail"]
    by_kpi = {r[1]: r for r in rows[4:]}
    assert by_kpi["Availability"][2] == pytest.approx(0.9995)
    assert by_kpi["Availability"][3] == "target 99.9%"
    assert by_kpi["Outages"][2] == 3
    assert by_kpi["Peak"][2] == pytest.approx(1.2)
    restricted = next(r for r in rows if r[0] == "APM Services")
    assert restricted[1].startswith("Not included — requires") and "apm.view" in restricted[1]
    failed = next(r for r in rows if r[0] == "Top Outages")
    assert failed[1].startswith("Not generated")
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref.startswith("A4:D")
    # Section names link to their sheets.
    link = next(c for c in ws["A"] if c.value == "Device Availability")
    assert link.hyperlink is not None and "01 Device Availability" in link.hyperlink.location


def test_xlsx_numbers_are_typed_and_text_stays_text():
    ws = load(render_xlsx(make_meta(), sections()))["01 Device Availability"]
    header_row = next(c.row for c in ws["A"] if c.value == "Device")
    hdr = [c.value for c in ws[header_row]][:6]
    assert hdr == ["Device", "IP", "Status", "Uptime", "Avg RTT", "Samples"]
    first = ws[header_row + 1]
    uptime, rtt, samples = first[3], first[4], first[5]
    assert uptime.value == pytest.approx(0.9995) and uptime.number_format == "0.00%"
    assert rtt.value == pytest.approx(12.5) and "ms" in rtt.number_format
    assert samples.value == 1234 and isinstance(samples.value, int)
    second = ws[header_row + 2]
    assert second[3].value == "—" and second[4].value == "—"     # unknown stays text
    third = ws[header_row + 3]
    assert third[3].value == pytest.approx(1.0) and third[3].number_format == "0.00%"
    assert third[5].value == "007"                               # leading zeros are identifiers
    assert first[1].value == "10.0.0.1"                          # mono column untouched
    # Freeze panes sit under the first table's header.
    assert ws.freeze_panes == f"A{header_row + 1}"
    # Every table with rows is an Excel table (filterable).
    assert ws.tables, "expected Excel tables"
    first_table = next(iter(ws.tables.values()))
    assert first_table.ref.startswith(f"A{header_row}:")
    assert first_table.autoFilter is not None
    # Chart series are exported as data.
    values = [v for row in cells(ws) for v in row]
    assert "Chart data — Trend" in values and "Time (GMT+3)" in values
    assert datetime(2026, 10, 1, 3, 0) in values


def test_xlsx_formula_injection_is_neutralised():
    ws = load(render_xlsx(make_meta(), sections()))["01 Device Availability"]
    texts = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    risky = [c for c in texts if c.value[:1] in ("=", "+", "-", "@", "\t", "\r")]
    assert not risky, [c.value for c in risky]
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
    devices = [c.value for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("'")]
    assert "'=HYPERLINK(\"http://evil\")" in devices
    assert "'@SUM(1+1)" in devices and "'+cmd" in devices and "'-2+3" in devices
    assert "'\tTabbed" in devices
    assert safe_text("=1+1") == "'=1+1" and safe_text("\rx") == "'\rx" and safe_text("ok") == "ok"
    assert safe_text("bad\x00\x07char") == "badchar"


def test_xlsx_header_uses_brand_colour_with_readable_text():
    light = make_branding(brand_color="#FFD400")
    ws = load(render_xlsx(make_meta(light), sections()))["01 Device Availability"]
    header = next(c for c in ws["A"] if c.value == "Device")
    assert header.fill.fgColor.rgb.endswith("FFD400")
    assert header.font.bold
    assert header.font.color.rgb.endswith("111827")          # dark text on a pale brand colour
    dark = load(render_xlsx(make_meta(make_branding(brand_color="#0B3D91")), sections()))
    header = next(c for c in dark["01 Device Availability"]["A"] if c.value == "Device")
    assert header.font.color.rgb.endswith("FFFFFF")


def test_xlsx_embeds_png_logo_but_not_svg():
    with_png = load(render_xlsx(make_meta(make_branding(logo=logo(png_bytes()))), sections()))
    assert len(with_png["Cover"]._images) == 1
    ext = with_png["Cover"]._images[0].anchor.ext       # displayed size, EMU (9525 per pixel)
    assert round(ext.height / 9525) == 56 and round(ext.width / 9525) == 168
    with_svg = load(render_xlsx(make_meta(make_branding(logo=logo(SVG))), sections()))
    assert len(with_svg["Cover"]._images) == 0


def test_xlsx_sheet_names_are_sanitised_unique_and_short():
    used: set[str] = set()
    a = sheet_name(1, "In/Out: [Top]* traffic? by \\ interface and very long name", used)
    b = sheet_name(1, "In/Out: [Top]* traffic? by \\ interface and very long name", used)
    assert len(a) <= 31 and len(b) <= 31 and a.lower() != b.lower()
    assert not re.search(r"[\[\]:*?/\\]", a + b)
    secs = [{"id": "x", "title": "Same", "tables": [{"headers": ["A"], "rows": [["1"]]}]},
            {"id": "y", "title": "Same", "tables": [{"headers": ["A"], "rows": [["2"]]}]}]
    wb = load(render_xlsx(make_meta(), secs))
    assert wb.sheetnames == ["Cover", "Summary", "01 Same", "02 Same"]


def test_xlsx_handles_empty_and_duplicate_headers_deterministically():
    secs = [{"id": "x", "title": "Edge", "tables": [
        {"headers": ["Name", "name", ""], "rows": [["a", "b", "c", "extra"]]},
        {"title": "Nothing", "headers": ["A", "B"], "rows": []}],
        "notes": ["note"]}]
    meta = make_meta()
    first = render_xlsx(meta, secs)
    second = render_xlsx(meta, secs)
    wb1, wb2 = load(first), load(second)
    for name in wb1.sheetnames:
        assert cells(wb1[name]) == cells(wb2[name])
    ws = wb1["01 Edge"]
    table = next(iter(ws.tables.values()))
    names = [c.name for c in table.tableColumns]
    assert len({n.lower() for n in names}) == len(names) == 4
    values = [v for row in cells(ws) for v in row]
    assert "No data for this window" in values and "note" in values


def test_parse_number_rules():
    assert parse_number("99.95%").value == pytest.approx(0.9995)
    assert parse_number("1,234").value == 1234
    assert parse_number("-3").value == -3
    assert parse_number("12.5 ms").unit == "ms"
    for text in ("—", "<0.01%", "5h 2m", "007", "10.0.0.1", "2026-10-07", "n/a"):
        assert parse_number(text) is None, text


# ─── Email summary ──────────────────────────────────────────────────────────

def many_sections(n: int = 40) -> list[dict]:
    out = []
    for i in range(n):
        out.append({"id": f"s{i}", "title": f"Section {i}",
                    "kpis": [{"label": f"Metric {i}", "value": f"{90 + i / 10:.1f}%",
                              "accent": ("danger", "warning", "success", "primary")[i % 4],
                              "subtitle": "vs target"}],
                    "tables": [{"headers": ["A", "B"], "rows": [["x" * 50, "y" * 50]] * 100}],
                    "notes": [f"Finding for section {i}"]})
    return out


def test_email_html_is_email_safe_and_branded():
    meta = make_meta(make_branding(logo=logo(png_bytes())))
    html = rs.email_summary_html(meta, sections(), "https://zen.example.com/api/v1/reports/shared/t?a=1&b=2",
                                 attached=True, logo_cid="report-logo@zenplus")
    assert "data:" not in html and "<svg" not in html.lower()
    assert "display:flex" not in html and "display:grid" not in html and "linear-gradient" not in html
    assert 'src="cid:report-logo@zenplus"' in html and 'alt="Acme Networks"' in html
    assert 'bgcolor="#0B3D91"' in html                       # brand header for Outlook
    assert 'width="600"' in html
    assert "CONFIDENTIAL" in html.upper() and "Confidential</td>" in html  # classification banner
    assert "Weekly Operations" in html and "GMT+3" in html
    assert "All monitored infrastructure" in html
    assert "Open full report" in html and "a=1&amp;b=2" in html
    assert "The full report is attached" in html
    assert "Acme NOC" in html and "https://acme.example" in html and "Internal use only." in html
    assert "Generated " in html
    # Top finding: the flagged figure with where to act on it, linked into the
    # appliance (the share link's origin); never the section's explanatory note.
    assert "Outages: 3. " in html and 'href="https://zen.example.com/availability"' in html
    assert "Review in Availability</a>" in html
    assert "Two devices breached the availability target." not in html
    assert "1 section not included" in html and "1 section could not be generated" in html
    # Restricted/failed sections do not leak KPIs; at most six KPIs.
    assert html.count("border-top:3px solid") <= 6


def test_email_html_without_logo_or_classification():
    meta = make_meta(make_branding(confidentiality=""))
    html = rs.email_summary_html(meta, sections(), None, attached=False)
    assert "cid:" not in html and "<img" not in html
    assert "Acme Networks" in html                             # name as text instead of a logo
    assert "letter-spacing:2px;text-transform:uppercase" not in html  # no banner
    assert "Open full report" not in html and "attached" not in html


def test_email_html_rejects_unsafe_view_url():
    html = rs.email_summary_html(make_meta(), sections(), "javascript:alert(1)", attached=False)
    assert "javascript:" not in html and "Open full report" not in html


def test_email_html_size_budget_with_40_sections():
    meta = make_meta(make_branding(logo=logo(png_bytes())))
    html = rs.email_summary_html(meta, many_sections(40), "https://zen.example.com/r/x", attached=True,
                                 logo_cid="report-logo@zenplus")
    assert len(html.encode("utf-8")) < 100 * 1024
    # One figure from each of the first three sections, each under its own caption.
    assert html.count("border-top:3px solid") == 3
    assert [t for t in re.findall(r'font-weight:700;color:#111827">(Section \d+)</div>', html)] == \
        ["Section 0", "Section 1", "Section 2"]
    findings = re.findall(r"<strong>Section (\d+)</strong>", html)
    assert len(findings) == 5
    # Danger findings first, the larger figure first within a severity
    # (sections 36, 32, ... read 93.6%, 93.2%, ...).
    assert findings[:5] == ["36", "32", "28", "24", "20"]
    assert "Finding for section" not in html                  # notes are never findings
    cols = re.findall(r'<td width="(\d+)%" valign="top" style="padding:5px">', html)
    assert set(cols) == {"50"}                                 # one figure per group → 2 columns


def test_email_text_alternative():
    meta = make_meta()
    text = rs.email_summary_text(meta, sections(), "https://zen.example.com/r/x")
    assert text.startswith("[CONFIDENTIAL]")
    assert "Acme Networks" in text and "Weekly Operations" in text
    assert "Period: " in text and "GMT+3" in text
    assert "- Availability: 99.95% (target 99.9%)" in text
    assert "Device Availability (All sources · over the report period)" in text
    assert "Top findings" in text and "- [!] Device Availability — Outages: 3. Review in Availability." in text
    linked = rs.email_summary_text(meta, sections(), "https://zen.example.com/api/v1/reports/shared/tok")
    assert "Outages: 3. Review in Availability: https://zen.example.com/availability" in linked
    assert "Two devices breached" not in text
    assert "Open the full report: https://zen.example.com/r/x" in text
    assert "Internal use only." in text and "Generated " in text
    assert "<" not in text.replace("<Co>", "")


def _noc_sections() -> list[dict]:
    """The field report's NOC briefing: current device status (device scope,
    snapshot) next to the fleet-wide open alert backlog."""
    return [
        {"id": "device_status_now", "title": "Current Device Status", "category": "Availability",
         "scope": "device", "semantics": "snapshot",
         "kpis": [{"label": "Network devices up", "value": "34", "accent": "success",
                   "subtitle": "of 34, at generation"},
                  {"label": "Down", "value": "0", "accent": "success"},
                  {"label": "Degraded", "value": "0", "accent": "neutral"},
                  {"label": "Unknown", "value": "0", "accent": "neutral"}],
         "findings": [], "notes": ["Every node in scope is up or not yet measured."]},
        {"id": "alert_backlog_aging", "title": "Open Alert Backlog", "category": "Alerting",
         "scope": "fleet", "semantics": "snapshot", "email_scope": "All alert sources",
         "kpis": [{"label": "Open alert incidents (all sources)", "value": "761", "accent": "primary"},
                  {"label": "Critical open incidents", "value": "285", "accent": "danger"},
                  {"label": "Unacknowledged", "value": "700", "accent": "warning"},
                  {"label": "Acknowledgement overdue", "value": "12", "accent": "danger"}],
         "findings": [{"severity": "danger", "text": "285 critical open incidents", "weight": 285,
                       "action": "Triage in Alert Center", "path": "/alerts?severity=critical"},
                      {"severity": "danger", "text": "12 incidents are past their acknowledgement target",
                       "weight": 12, "action": "Assign and acknowledge in Alert Center", "path": "/alerts"}],
         "notes": ["Acknowledgement targets (Alert Center settings): critical 15 min."]},
        {"id": "server_telemetry_coverage", "title": "Server Telemetry Coverage", "category": "Servers",
         "scope": "fleet", "semantics": "window",
         "kpis": [{"label": "Telemetry samples received", "value": "66.14%", "accent": "warning"}],
         "notes": ["Coverage compares the CPU samples stored for each agent-monitored server ..."]},
    ]


def test_email_groups_kpis_per_section_with_scope_and_as_of():
    meta = make_meta(scope_label="All monitored infrastructure")
    groups = rs.email_kpi_groups(_noc_sections(), meta)
    assert [g["title"] for g in groups] == ["Current Device Status", "Open Alert Backlog"]
    assert [len(g["kpis"]) for g in groups] == [3, 3]           # capped per section, six in all
    assert groups[0]["caption"] == f"All network devices · as of {meta['generated_label']}"
    assert groups[1]["caption"] == f"All alert sources · as of {meta['generated_label']}"
    # The cap keeps a section's alarm: headline, then danger, then warning figures.
    assert [k["label"] for k in groups[1]["kpis"]] == [
        "Open alert incidents (all sources)", "Critical open incidents", "Acknowledgement overdue"]
    # A device filter is named; a period section says so.
    scoped = rs.email_kpi_groups(_noc_sections(), {**meta, "scope_label": "3 selected device(s)"})
    assert scoped[0]["caption"].startswith("3 selected device(s) · as of ")
    window = rs.email_kpi_groups(_noc_sections()[2:], meta)
    assert window[0]["caption"] == "All sources · over the report period"

    html = rs.email_summary_html(meta, _noc_sections(), "https://nms.acme.example/api/v1/reports/shared/tok",
                                 attached=False)
    assert html.count("border-top:3px solid") == 6
    card = html[html.index('width="600"'):]                    # past the hidden preheader
    assert card.index("Current Device Status") < card.index("Network devices up") < \
        card.index("Open Alert Backlog") < card.index("Open alert incidents (all sources)")
    assert "All alert sources · as of " in html


def test_email_findings_are_explicit_actions_ranked_by_severity_and_size():
    found = rs._email_findings(_noc_sections())
    assert [(f["severity"], f["section"], f["text"]) for f in found] == [
        ("danger", "Open Alert Backlog", "285 critical open incidents"),
        ("danger", "Open Alert Backlog", "12 incidents are past their acknowledgement target"),
        ("warning", "Server Telemetry Coverage", "Telemetry samples received: 66.14%"),
    ]
    assert found[0]["path"] == "/alerts?severity=critical" and found[2]["path"] == "/servers"
    # A section that judged its figures (findings: []) adds none; notes never become findings.
    assert all(f["section"] != "Current Device Status" for f in found)
    assert not any("Acknowledgement targets" in f["text"] or "compares the CPU samples" in f["text"]
                   for f in found)
    html = rs.email_summary_html(make_meta(), _noc_sections(),
                                 "https://nms.acme.example/api/v1/reports/shared/tok", attached=False)
    assert 'href="https://nms.acme.example/alerts?severity=critical"' in html
    assert ">Triage in Alert Center</a>" in html
    # Without an appliance address the action is plain text.
    text = rs.email_summary_text(make_meta(), _noc_sections(), None)
    assert "- [!] Open Alert Backlog — 285 critical open incidents. Triage in Alert Center." in text
    # An unsafe path is never linked.
    evil = [{**_noc_sections()[1], "findings": [{"severity": "danger", "text": "x", "path": "//evil.example"}]}]
    assert "evil.example" not in rs.email_summary_html(make_meta(), evil, "https://nms.acme.example/api/v1/"
                                                       "reports/shared/t", attached=False)


def test_email_labels_manual_runs_and_the_schedule_time():
    meta = make_meta(run_kind="manual", run_note="Manual run: generated on request, not at the schedule's time.",
                     schedule_label="Daily at 08:00 Asia/Riyadh (05:00 UTC)")
    html = rs.email_summary_html(meta, sections(), None, attached=False)
    assert "Manual run: generated on request" in html
    assert "Schedule: Daily at 08:00 Asia/Riyadh (05:00 UTC) · this was a manual run" in html
    text = rs.email_summary_text(meta, sections(), None)
    assert "Manual run: generated on request" in text
    assert "Schedule: Daily at 08:00 Asia/Riyadh (05:00 UTC) · this was a manual run" in text
    scheduled = rs.email_summary_html(make_meta(run_kind="scheduled",
                                                schedule_label="Daily at 08:00 Asia/Riyadh (05:00 UTC)"),
                                      sections(), None, attached=False)
    assert "Scheduled: Daily at 08:00 Asia/Riyadh (05:00 UTC)</div>" in scheduled
    assert "Manual run" not in scheduled


def test_email_logo_attachment_png_jpeg_only():
    png = png_bytes()
    att = rs.email_logo_attachment(make_meta(make_branding(logo=logo(png))))
    assert att == ("report-logo@zenplus", png, "image/png")

    buf = io.BytesIO()
    Image.new("RGB", (20, 10), (0, 0, 0)).save(buf, "JPEG")
    jpeg = buf.getvalue()
    att = rs.email_logo_attachment(make_meta(make_branding(logo=logo(jpeg))))
    assert att is not None and att[2] == "image/jpeg"

    assert rs.email_logo_attachment(make_meta(make_branding(logo=logo(SVG)))) is None
    assert rs.email_logo_attachment(make_meta(make_branding())) is None
    # SVG light logo + PNG dark logo → the dark PNG (on a dark cell).
    meta = make_meta(make_branding(logo=logo(SVG), logo_dark=logo(png)))
    att = rs.email_logo_attachment(meta)
    assert att is not None and att[1] == png
    html = rs.email_summary_html(meta, sections(), None, attached=False, logo_cid=att[0])
    assert 'bgcolor="#1F2937"' in html


# ─── Legacy HTML email / page ───────────────────────────────────────────────

def legacy_meta(**over) -> dict:
    meta = {"title": "Weekly Executive Summary", "report_type": "executive_summary",
            "period": "last_7d", "period_label": "Last 7 days",
            "generated_at": datetime(2026, 10, 7, 5, 0, tzinfo=timezone.utc),
            "company_name": "Acme Networks", "product_name": "ZenPlus"}
    meta.update(over)
    return meta


LEGACY_DATA = {"kpis": {"availability_pct": 99.2, "incidents_count": 3, "mttr_minutes": 12,
                        "devices_monitored": 42, "active_critical_count": 1,
                        "sla_attained_pct": 99.2, "sla_target_pct": 99.9}}


def test_legacy_email_uses_company_branding_not_product_badge():
    b = make_branding(brand_color="#7A1FA2")
    html = report_html.build_report_email_html(
        legacy_meta(branding=b, start=FRM, end=TO), LEGACY_DATA,
        view_url="https://zen.example.com/r/1", attached=True)
    assert ">Z</div>" not in html
    assert "Acme Networks" in html
    assert "#7A1FA2" in html and 'bgcolor="#7A1FA2"' in html
    assert "2026-10-07 08:00 GMT+3" in html                    # generated, org timezone
    assert "Last 7 days (2026-10-01 03:00 — 2026-10-08 03:00 GMT+3)" in html
    assert "Internal use only." in html and "Acme NOC" in html
    assert "View full report" in html and "SCHEDULED REPORT" in html


def test_legacy_page_without_branding_falls_back_to_utc_label():
    html = report_html.build_report_html(legacy_meta(), LEGACY_DATA)
    assert "Acme Networks" in html and ">Z</div>" not in html
    assert "Oct 07, 2026 · 05:00 UTC" in html
    text = report_html.build_report_email_text(legacy_meta(branding=make_branding()), LEGACY_DATA)
    assert "Acme Networks · Last 7 days" in text and "GMT+3" in text
    assert "— Acme Networks automated report" in text


# ─── Legacy Excel / CSV export ──────────────────────────────────────────────

FLEET_ALERT = {"id": 1, "device_id": "11111111-1111-1111-1111-111111111111", "service_check_id": None,
               "status": "active", "severity": "critical", "message": "=HYPERLINK(\"http://evil\",\"x\")",
               "triggered_at": datetime(2026, 10, 7, 10, 0, tzinfo=timezone.utc),
               "acknowledged_at": None, "resolved_at": None, "hostname": "other-team-router"}


@pytest.fixture
def fake_export(monkeypatch):
    calls: dict[str, list] = {"alerts": [], "ping": [], "status": []}
    visible: list[dict] = []

    async def fetch_devices(db, device_ids=None, group_ids=None, locations=None, device_types=None):
        if device_ids == [scoping.NO_MATCH_ID]:
            return []
        return [d for d in visible if not device_ids or d["id"] in device_ids]

    async def fetch_alerts(db, start, end, device_ids=None, visible_tags=None, *, redact_netflow=False):
        calls["alerts"].append(device_ids)
        if not device_ids:
            return [FLEET_ALERT]  # the real helper treats "no ids" as "every device"
        return [dict(FLEET_ALERT, device_id=i, hostname="mine") for i in device_ids]

    def fetch_ping(start, end, device_ids=None):
        calls["ping"].append(device_ids)
        return []

    def fetch_status(start, end, device_ids=None):
        calls["status"].append(device_ids)
        return [] if device_ids else [{"device_id": "x", "timestamp": TO, "old_status": "up",
                                       "new_status": "down", "reason": "leak", "duration_sec": 1}]

    monkeypatch.setattr(export_service, "_fetch_devices", fetch_devices)
    monkeypatch.setattr(export_service, "_fetch_alerts", fetch_alerts)
    monkeypatch.setattr(export_service, "_fetch_ping_metrics", fetch_ping)
    monkeypatch.setattr(export_service, "_fetch_device_status_log", fetch_status)

    async def no_services(db, visible_tags=None):
        raise AssertionError("services are not part of this export")

    monkeypatch.setattr(export_service, "_fetch_service_checks", no_services)
    return calls, visible


@pytest.mark.parametrize("report_type", ["alert_analysis", "executive_summary", "device_health"])
def test_scoped_account_with_no_visible_devices_gets_no_fleet_data(fake_export, report_type):
    calls, _ = fake_export
    kwargs = dict(db=object(), report_type=report_type, period="last_24h",
                  device_ids=[scoping.NO_MATCH_ID], branding=make_branding())
    xlsx = asyncio.run(export_service.generate_excel_report(**kwargs))
    body = asyncio.run(export_service.generate_csv_report(**kwargs)).decode()
    assert calls["alerts"] == [] and calls["ping"] == [] and calls["status"] == []
    wb = load(xlsx)
    everything = " ".join(str(v) for ws in wb.worksheets for row in cells(ws) for v in row if v is not None)
    assert "other-team-router" not in everything and "HYPERLINK" not in everything
    assert "leak" not in everything
    assert "other-team-router" not in body


def test_export_with_visible_devices_is_confined_to_them(fake_export):
    calls, visible = fake_export
    visible.append({"id": "22222222-2222-2222-2222-222222222222", "hostname": "mine", "status": "up",
                    "ip_address": "10.0.0.9", "device_type": "router", "group_name": "Core",
                    "location": "HQ", "last_seen": datetime(2026, 10, 7, 9, 30, tzinfo=timezone.utc)})
    body = asyncio.run(export_service.generate_csv_report(
        db=object(), report_type="alert_analysis", period="last_24h",
        device_ids=["22222222-2222-2222-2222-222222222222"], branding=make_branding()))
    assert calls["alerts"] == [["22222222-2222-2222-2222-222222222222"]]
    rows = list(csv.reader(io.StringIO(body.decode())))
    assert rows[0][0] == "Organisation: Acme Networks"
    assert any(r and r[0].startswith("Period: Last 24 Hours (") and "GMT+3" in r[0] for r in rows)
    header = next(r for r in rows if r and r[0].startswith("Triggered At"))
    assert header[0] == "Triggered At (GMT+3)" and header[5] == "Acknowledged At (GMT+3)"
    data = rows[rows.index(header) + 1]
    assert data[0] == "2026-10-07 13:00:00"                    # 10:00 UTC in Asia/Riyadh
    assert data[4].startswith("'=HYPERLINK")                   # CSV injection neutralised


def test_excel_export_is_branded_with_cover_and_local_times(fake_export):
    _, visible = fake_export
    visible.append({"id": "22222222-2222-2222-2222-222222222222", "hostname": "=evil()", "status": "up",
                    "ip_address": "10.0.0.9", "device_type": "router", "group_name": "Core",
                    "location": "HQ", "last_seen": datetime(2026, 10, 7, 9, 30, tzinfo=timezone.utc)})
    wb = load(asyncio.run(export_service.generate_excel_report(
        db=object(), report_type="device_health", period="last_24h",
        device_ids=["22222222-2222-2222-2222-222222222222"],
        branding=make_branding(brand_color="#0B3D91", date_format="DD/MM/YYYY"))))
    assert wb.sheetnames[0] == "Cover"
    cover = [v for row in cells(wb["Cover"]) for v in row if v is not None]
    assert "Acme Networks" in cover and "Device Health Report" in cover
    assert any(isinstance(v, str) and v.startswith("Last 24 Hours (") and "GMT+3" in v for v in cover)
    ws = wb["Device Overview"]
    header = [c.value for c in ws[1]]
    assert header[-1] == "Last Seen (GMT+3)"
    assert ws.cell(row=1, column=1).fill.fgColor.rgb.endswith("0B3D91")
    last_seen = ws.cell(row=2, column=len(header))
    assert last_seen.value == datetime(2026, 10, 7, 12, 30)
    assert last_seen.number_format.startswith("dd/mm/yyyy")
    assert ws.cell(row=2, column=1).value == "'=evil()"
    assert wb.properties.creator == "Acme Networks"
